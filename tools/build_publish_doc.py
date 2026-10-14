"""生成《个人开发者鸿蒙应用上架全流程》DOCX（内容与 docs/store/RELEASE_RUNBOOK.md 一致）。

用法： <python> tools/build_publish_doc.py
输出： C:\\Users\\daode\\OneDrive\\Desktop\\个人开发者鸿蒙应用上架全流程.docx
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Microsoft YaHei'
BRAND = (0xB3, 0x6B, 0x14)
DARK = (0x3A, 0x32, 0x2A)
GREY = (0x6F, 0x67, 0x5C)

OUT = os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop',
                   '个人开发者鸿蒙应用上架全流程.docx')
if not os.path.isdir(os.path.dirname(OUT)):
    OUT = os.path.join(os.path.expanduser('~'), 'Desktop',
                       '个人开发者鸿蒙应用上架全流程.docx')


def set_font(run, size=10.5, bold=False, color=DARK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), FONT)


def set_style_font(style, size, bold=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)


def para(doc, text, size=10.5, bold=False, color=DARK, space_after=6, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    set_font(p.add_run(text), size, bold, color)
    return p


def bullet(doc, text, size=10.5, color=DARK):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), size, False, color)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(htext), 9.5, True, BRAND)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run(val), 9, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def main():
    doc = Document()
    set_style_font(doc.styles['Normal'], 10.5, False, DARK)
    for name, size in (('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11)):
        set_style_font(doc.styles[name], size, True, BRAND)

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # 标题
    title = doc.add_paragraph()
    title.alignment = 1
    set_font(title.add_run('个人开发者鸿蒙应用上架全流程'), 20, True, BRAND)
    sub = doc.add_paragraph()
    sub.alignment = 1
    sub.paragraph_format.space_after = Pt(16)
    set_font(sub.add_run('以「心情手记」为例 · HarmonyOS / AppGallery Connect'), 11, False, GREY)

    # 0
    doc.add_heading('0. 当前的完成情况', level=1)
    para(doc, '工程侧改造已全部完成，剩下的是账号侧操作。先看两张表，知道自己站在哪一步。', color=GREY)
    para(doc, '已完成（无需再做）', bold=True)
    table(doc, ['项', '值'], [
        ['应用名', '心情手记'],
        ['包名', 'com.datelog.couple（保持不变；包名不参与审核、用户不可见）'],
        ['版本', '1.4.0 / versionCode 1000006'],
        ['签名材料', 'C:\\Users\\daode\\.ohos\\config\\ 下 datelog.p12 / datelog-release.cer / datelog-release.p7b'],
        ['签名口令', '<签名口令：见本机私密文件 tools\signing.local.txt>，别名 datelog（发布证书有效期至 2029-10-08）'],
        ['构建配置', '签名口令已由 DevEco 加密写入 build-profile.json5，可直接命令行出签名包'],
        ['内置头像', '6 张（男 3 / 女 3），不读取相册'],
        ['隐私政策', 'v2.3；零网络权限、零数据采集'],
        ['图标', 'docs/store/app-icon-square-1024.png（直角方图，AGC 要的就是这个）'],
    ], [3.4, 12.6])

    para(doc, '待完成（本次上架必须做的四件事）', bold=True)
    table(doc, ['#', '事项', '说明'], [
        ['1', '覆盖上传隐私政策到 OBS', 'v2.3 改了「12 张头像 → 6 张」与「心情支持自定义」，URL 不变'],
        ['2', '重拍 5 张截图', '必须是 9:16 或 16:9；AGC 默认选「横向」，必须切「竖向截图」'],
        ['3', 'AGC 改两项', '应用名 →「心情手记」；分类与主标签 →「工具」'],
        ['4', '真机回归 4 点', '首页顶部空白 / 跨零点封存 / 补记不污染今日状态 / 6 宫格头像'],
    ], [1.0, 4.6, 10.4])

    doc.add_page_break()

    # 1
    doc.add_heading('1. 前置：实名认证', level=1)
    para(doc, '打开 AppGallery Connect（developer.huawei.com/consumer/cn/service/josp/agc/index.html）'
              '登录华为开发者账号，完成实名认证（个人：身份证）。', indent=0.4)
    para(doc, '未实名无法创建应用、也无法申请发布证书。已有实名账号并已建好本应用项目的，直接跳到第 2 节。',
         color=GREY, indent=0.4)

    # 2
    doc.add_heading('2. AGC 项目与应用', level=1)
    para(doc, '1. AGC 控制台 →「我的项目」→ 进入本应用所在项目。', indent=0.4)
    para(doc, '2. 若尚未创建应用：「添加应用」→ 平台 HarmonyOS → 包名 com.datelog.couple，'
              '必须与 AppScope/app.json5 的 bundleName 完全一致。', indent=0.4)
    para(doc, '3. 本次必改两项（上次被拒的直接原因）：', bold=True, indent=0.4)
    table(doc, ['位置', '原值', '改为'], [
        ['应用名称', 'DateLog', '心情手记'],
        ['应用分类', '生活服务', '工具'],
        ['主标签', '（未设或非工具）', '工具'],
    ], [3.2, 5.2, 7.6])
    para(doc, '依据：上次驳回原文「分类/主标签与实际功能不符，应提交至【工具】分类并设置【工具】为主标签」，'
              '对应审核指南 3.5 —— 本应用是单机记录工具，不属于社交或生活服务类。', color=GREY, indent=0.4)

    # 3
    doc.add_heading('3. 签名（材料已有，无需重做）', level=1)
    para(doc, '当前签名可直接使用，不需要重新申请证书或 Profile。', indent=0.4)
    table(doc, ['项', '值'], [
        ['位置', 'C:\\Users\\daode\\.ohos\\config\\'],
        ['文件', 'datelog.p12、datelog-release.cer、datelog-release.p7b'],
        ['口令 / 别名', '<签名口令：见本机私密文件 tools\signing.local.txt> / datelog'],
        ['有效期', '发布证书至 2029-10-08'],
    ], [3.4, 12.6])
    para(doc, '注意：.p12 与口令绝不能丢。丢失后无法再签同一个应用，只能换包名重新上架。'
              '桌面「DateLog-签名文件」有一份备份，建议再存一份到网盘或 U 盘。', bold=True, indent=0.4)
    para(doc, '若将来必须重新配置签名，只能走 DevEco Studio：File → Project Structure → '
              'Signing Configs，取消自动签名后手动指定 p12 / p7b / cer。', indent=0.4)
    para(doc, '实测坑：手填明文口令会构建失败，报 00303116「The length of the storePassword or '
              'keyPassword field ... is less than 32」。hvigor 只接受 DevEco 加密后的口令串'
              '（AES-GCM，密文必然不小于 32 字节），而 decipher-util 只提供 decryptPwd()、'
              '没有加密方法 —— 加密逻辑只在 IDE 内部。', color=GREY, indent=0.4)

    doc.add_page_break()

    # 4
    doc.add_heading('4. 构建与自检', level=1)
    doc.add_heading('4.1 命令行出签名包', level=2)
    for line in [
        '$env:PATH = "E:\\huawei\\DevEco Studio\\tools\\node;E:\\huawei\\DevEco Studio\\tools\\ohpm\\bin;'
        'E:\\huawei\\DevEco Studio\\tools\\hvigor\\bin;$env:PATH"',
        '$env:DEVECO_SDK_HOME = "E:\\huawei\\DevEco Studio\\sdk"',
        'hvigorw assembleApp --mode project -p product=default -p buildMode=release --no-daemon',
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(line), 9, False, (0x20, 0x30, 0x50))
    para(doc, '产物：build\\outputs\\default\\DevecoStudioProject-default-signed.app', indent=0.4)
    para(doc, '注意：hvigor 把日志写到 stderr，PowerShell 里即使 BUILD SUCCESSFUL 也可能显示 '
              'exit code: 1。看输出文字，不要看退出码。', color=GREY, indent=0.4)

    doc.add_heading('4.2 上传前自检', level=2)
    for item in [
        '产物是 *-signed.app（不是 *-unsigned.app）',
        '官方验签 hap-sign-tool verify-app 返回 Verify success，退出码 0',
        'bundleName = com.datelog.couple',
        'versionCode / versionName 已确认（当前 1000006 / 1.4.0）',
        '单元测试全绿：hvigorw test --mode module -p product=default --no-daemon（当前 47/47）',
        "构建日志没有 “Will skip sign 'app'”（出现说明签名没生效）",
    ]:
        bullet(doc, item)
    para(doc, '重要：不要把签名包解压再打包。签名块存在于 ZIP 中央目录之外，用 Expand-Archive 解压会丢签名，'
              '之后验签会误报「signature not found」。始终校验原始产物。', bold=True, indent=0.4)

    doc.add_heading('4.3 真机回归（本次改动相关）', level=2)
    table(doc, ['#', '验证点', '期望'], [
        ['1', '首页顶部空白', '确认空白消失（已改为 Scroll 满高 + 顶对齐）'],
        ['2', '跨零点封存', '今天写完 → 系统日期调到明天 → 重开该记录 →「编辑」消失'],
        ['3', '补记不污染今日状态', '从日历补记某天 → 保存后首页仍显示「今天未记录」'],
        ['4', '头像选择', '「我的」→ 点头像 → 6 宫格（3x2）能正常选中与切换'],
    ], [1.0, 4.4, 10.6])
    para(doc, '建议测试路径（也用于 AGC 审核备注）：', bold=True, indent=0.4)
    para(doc, '同意隐私政策 → 填写称呼 → 首页「记录今天」→ 选心情 + 写正文与明日计划 → 保存 → '
              '首页「最近的记录」点开全文 → 时间轴点一条看详情 → 点「编辑」改完返回 → 日历点某天 → '
              '「我的」切换日夜模式、更换头像、进入任务规划新建并完成一个任务 → 桌面添加服务卡片。',
         color=GREY, indent=0.4)

    doc.add_page_break()

    # 5
    doc.add_heading('5. 商店资料', level=1)

    doc.add_heading('5.1 图标', level=2)
    bullet(doc, '必须上传直角方图：docs/store/app-icon-square-1024.png')
    bullet(doc, '不要传任何已做圆角的预览图 —— AGC 会判不合格')

    doc.add_heading('5.2 截图（3 到 5 张）', level=2)
    bullet(doc, '尺寸只能是 9:16 或 16:9。模拟器原始截图（1256x2760 约等于 9:19.8）不合格，需合成到 1080x1920。')
    bullet(doc, 'AGC 上传页默认选中「横向截图」，必须手动切到「竖向截图」，否则尺寸校验不过。')
    bullet(doc, '合成脚本：tools/generate_store_screenshots.py。脚本里的文件名前缀目前是占位符，'
                '重拍后必须改成实际文件名前缀，否则脚本直接报错（有意设计）。')

    doc.add_heading('5.3 隐私政策', level=2)
    bullet(doc, '把 docs/store/privacy-policy.html 上传为 OBS 桶 datelog-privacy 的 index.html（覆盖）。')
    bullet(doc, '公网 URL：https://datelog-privacy.obs.cn-north-4.myhuaweicloud.com/index.html')
    bullet(doc, 'AGC 只接受公网可访问地址，不接受本地文件路径。')
    bullet(doc, '验证：用无痕窗口打开该 URL，应直接显示政策正文、无需登录。')
    bullet(doc, '上传指南见 docs/store/OBS_HOSTING_GUIDE.md。')

    doc.add_heading('5.4 文案', level=2)
    para(doc, '打开 docs/store/LISTING.md，把「一句话简介」「详细描述」「权限说明」「内容分级」'
              '「审核备注」「版本更新说明」逐段复制进 AGC 表单。', indent=0.4)

    doc.add_heading('5.5 权限说明（AGC 权限页）', level=2)
    table(doc, ['权限', '处理'], [
        ['定位服务', 'AGC 默认勾选 → 必须取消（本应用不使用）'],
        ['华为账号', 'AGC 默认勾选 → 必须取消（本应用不使用）'],
        ['代理提醒', '保留勾选（「重要日子提醒」确实使用）'],
        ['网络权限', '本应用未申请，无需说明'],
    ], [3.4, 12.6])

    # 6
    doc.add_heading('6. 提交审核', level=1)
    para(doc, '1. AGC → 你的应用 →「版本管理」→ 新建版本。', indent=0.4)
    para(doc, '2. 上传第 4 节得到的 *-signed.app。', indent=0.4)
    para(doc, '3. 填写：版本号 1.4.0、更新说明、隐私政策 URL、应用分类「工具」、截图、内容分级、权限说明。', indent=0.4)
    para(doc, '4. 提交上架审核。', indent=0.4)
    para(doc, '5. 通过后选择「手动发布」或「审核通过后自动发布」。', indent=0.4)
    para(doc, '审核时长参考：预审 30 到 40 分钟，人工审核数小时至 1 到 3 个工作日。', color=GREY, indent=0.4)

    doc.add_page_break()

    # 7
    doc.add_heading('7. 上次被拒的三条问题与对应整改', level=1)
    table(doc, ['上次驳回', '根因', '本次整改'], [
        ['应用信息：分类/主标签与实际功能不符，应提交至【工具】', 'AGC 里选的是「生活服务」',
         '分类与主标签都改为「工具」'],
        ['应用功能：仅有简单记录打卡功能，使用场景有限（审核指南 3.5）', '功能面偏薄',
         '已扩展为完整记录工具：心情（含自定义心情）+ 手写正文 + 明日计划回显 + 记录详情与回看 + 历史封存 + '
         '任务规划（最多 6 个）+ 日历/时间轴/统计 + 图片记录 + 重要日子 + 桌面卡片 + 日夜模式 + 6 张内置头像'],
        ['真机（nova 14 Pro / HarmonyOS 6.1.0）修改头像无反应、添加照片无反应',
         'fileIo.copyFile 传媒体库 URI 会抛 13900020',
         '已改为 openSync(uri) → openSync(dest) → copyFile(fd, fd)，并移除 3 处静默失败路径。'
         '本轮头像彻底改为内置头像，不再读取相册，该问题从根上消失'],
    ], [4.6, 4.4, 7.0])

    # 附录
    doc.add_heading('附录 A · 常见驳回原因与对策', level=1)
    table(doc, ['驳回原因', '对策'], [
        ['上传了未签名包', '确认产物是 *-signed.app'],
        ['包名与 AGC 项目不一致', '必须是 com.datelog.couple'],
        ['缺隐私政策 / URL 打不开', 'OBS 覆盖上传后，用无痕窗口确认可直接打开'],
        ['未提供首启隐私同意', '已内置（首启展示政策 + 同意/退出）'],
        ['权限说明缺失或不实', 'AGC 权限页取消「定位」「华为账号」，保留「代理提醒」'],
        ['分类/主标签不符', '都设为「工具」'],
        ['截图尺寸不合格', '只能 9:16 或 16:9，且要切「竖向截图」'],
        ['图标不合格', '传直角方图，不要传已圆角的预览图'],
        ['截图与实机不符 / 全空状态', '先造数据再截图'],
        ['versionCode 未递增', '每次提交必须比上一版大'],
    ], [5.4, 10.6])

    doc.add_heading('附录 B · 版本迭代流程', level=1)
    for item in [
        '改 AppScope/app.json5：versionCode 加一，按需改 versionName',
        'hvigorw assembleApp --mode project -p product=default -p buildMode=release --no-daemon',
        '官方验签 + 真机回归',
        'AGC 新建版本 → 上传新 *-signed.app → 填更新说明 → 提交',
        '发布证书与 .p12 长期复用，无需重新申请',
    ]:
        bullet(doc, item)

    doc.add_heading('附录 C · 可选加固（非必需）', level=1)
    bullet(doc, '代码混淆：当前 release 配置里 obfuscation.enable = false。上架不强制混淆；'
                '若要开启，先改成 true，再真机完整回归（-enable-filename-obfuscation 可能影响 '
                'router 页面路径与卡片 LocalStorage 键）。')
    bullet(doc, '降低系统版本门槛：当前 compatibleSdkVersion 为 6.1.1(24)。想覆盖更老设备需装对应 SDK、'
                '下调该值并全量回归，属独立评估项。')
    bullet(doc, '头像与包体：6 张头像每张约 50 KB（柔和渐变插画）。若要瘦身可量化成 256 色 PNG'
                '（总量约能降到 100 KB），但要重新目视确认无渐变色带。')

    doc.add_heading('附录 D · 环境与踩坑备忘', level=1)
    table(doc, ['项', '值 / 说明'], [
        ['DevEco Studio', 'E:\\huawei\\DevEco Studio'],
        ['SDK', 'E:\\huawei\\DevEco Studio\\sdk（DEVECO_SDK_HOME）'],
        ['工程根', 'E:\\DevecoStudioProject'],
        ['Python（脚本用）', 'C:\\Users\\daode\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe'],
        ['真实桌面', 'C:\\Users\\daode\\OneDrive\\Desktop（不是 C:\\Users\\daode\\Desktop，后者是残留）'],
        ['hvigor 退出码', '日志走 stderr，成功也可能报 exit code: 1 —— 看文字'],
        ['AGC 下载按钮', '要左键点「下载」，右键「另存为网页」会存成 HTML 而不是证书'],
        ['桌面 DateLog-* 命名', '历史遗留命名，内容已更新为「心情手记」'],
    ], [3.6, 12.4])

    doc.save(OUT)
    print('已生成：%s (%.1f KB)' % (OUT, os.path.getsize(OUT) / 1024.0))


if __name__ == '__main__':
    main()
